"""
ppt_chart_toolkit.py — python-pptx 图表封装（R106, 2026-03-26）

提供 add_chart(slide, chart_type, data, l, t, w, h, title) 统一接口。
chart_type: "COLUMN" | "BAR" | "LINE" | "PIE"
data: {"categories": [...], "series": {"系列名": [值, ...]}}
l, t, w, h: 英寸，左/上/宽/高
返回 Chart 对象

依赖: pip install python-pptx（已安装）
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE

_CHART_TYPE_MAP = {
    "COLUMN": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "BAR":    XL_CHART_TYPE.BAR_CLUSTERED,
    "LINE":   XL_CHART_TYPE.LINE,
    "PIE":    XL_CHART_TYPE.PIE,
}


def add_chart(slide, chart_type, data, l=1.0, t=1.5, w=8.0, h=4.5, title=None):
    """
    在 python-pptx slide 对象上插入图表。
    chart_type: "COLUMN" | "BAR" | "LINE" | "PIE"
    data: {"categories": [str,...], "series": {"名": [num,...]}}
    l, t, w, h: 英寸
    title: 图表标题字符串，None则不显示
    返回 Chart 对象
    """
    xl_type = _CHART_TYPE_MAP.get(chart_type.upper(), XL_CHART_TYPE.COLUMN_CLUSTERED)
    cd = ChartData()
    cd.categories = data["categories"]
    for series_name, values in data["series"].items():
        cd.add_series(series_name, values)
    chart_shape = slide.shapes.add_chart(
        xl_type,
        Inches(l), Inches(t), Inches(w), Inches(h),
        cd
    )
    chart = chart_shape.chart
    if title:
        chart.has_title = True
        chart.chart_title.text_frame.text = title
    else:
        chart.has_title = False
    return chart


def new_prs():
    """新建空白 python-pptx 演示文稿，返回 Presentation 对象"""
    return Presentation()


def add_blank_slide(prs):
    """添加空白幻灯片（layout 6），返回 slide 对象"""
    blank_layout = prs.slide_layouts[6]
    return prs.slides.add_slide(blank_layout)


def save_prs(prs, path):
    """保存 Presentation 到指定路径"""
    prs.save(path)
    return path


# ── 便捷常量 ──────────────────────────────────────────────
CHART_COLUMN = "COLUMN"
CHART_BAR    = "BAR"
CHART_LINE   = "LINE"
CHART_PIE    = "PIE"
