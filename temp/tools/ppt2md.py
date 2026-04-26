#!/usr/bin/env python3
"""
ppt2md.py - PPT幻灯片 -> 结构化Markdown讲义导出工具
用法:
  python ppt2md.py <input.pptx> [output.md]
  python ppt2md.py <input.pptx> --stdout

功能:
  - 提取每页幻灯片标题和正文要点
  - 识别封面页/目录页/内容页
  - 输出结构化Markdown，含课程元数据、知识点列表
"""
import sys, os, re
from pptx import Presentation
from pptx.util import Pt


def extract_slide_texts(slide):
    """提取单张幻灯片所有文本，按shape顺序，返回 list of (level, text)"""
    items = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            # 过滤纯emoji、纯数字页码
            clean = re.sub(r"[\U0001F000-\U0001FFFF\u2600-\u27FF]", "", text).strip()
            if not clean or re.fullmatch(r"\d{1,2}", clean):
                continue
            items.append((para.level, text))
    return items


def classify_slide(idx, texts):
    """判断幻灯片类型: cover/toc/overview/content"""
    if idx == 0:
        return "cover"
    flat = " ".join(t for _, t in texts)
    if any(kw in flat for kw in ["教学大纲", "目录", "大纲", "课程结构"]):
        return "toc"
    if any(kw in flat for kw in ["教学流程", "课程概览", "流程总览", "数据概览"]):
        return "overview"
    return "content"

def ppt_to_markdown(pptx_path: str) -> str:
    """读取PPT，返回结构化Markdown字符串"""
    prs = Presentation(pptx_path)
    total = len(prs.slides)
    md_lines = []

    # 元数据头
    fname = os.path.basename(pptx_path)
    md_lines.append(f"# {fname}")
    md_lines.append(f"")
    md_lines.append(f"> 共 {total} 页幻灯片  ")
    md_lines.append(f"> 自动导出自 ppt2md.py")
    md_lines.append("")

    knowledge_points = []  # 收集所有知识点
    section_titles = []    # 章节标题

    for idx, slide in enumerate(prs.slides):
        texts = extract_slide_texts(slide)
        if not texts:
            continue
        kind = classify_slide(idx, texts)

        if kind == "cover":
            # 封面：提取课程标题
            md_lines.append("## 课程信息")
            md_lines.append("")
            for lvl, t in texts:
                md_lines.append(f"**{t}**")
            md_lines.append("")
            continue

        if kind in ("toc", "overview"):
            continue  # 跳过目录/流程总览页

        # 内容页处理
        page_num = idx + 1
        if not texts:
            continue

        # 第一条非空文本作为页标题
        title_text = texts[0][1] if texts else f"第{page_num}页"
        # 去除emoji
        title_clean = re.sub(r"[\U0001F000-\U0001FFFF\u2600-\u27FF\U0001F300-\U0001F9FF]", "", title_text).strip()

        md_lines.append(f"## [{page_num}] {title_clean}")
        md_lines.append("")

        if title_clean not in section_titles:
            section_titles.append(title_clean)

        # 正文要点
        body_items = texts[1:]
        for lvl, t in body_items:
            t_clean = re.sub(r"[\U0001F000-\U0001FFFF\u2600-\u27FF\U0001F300-\U0001F9FF]", "", t).strip()
            if not t_clean:
                continue
            indent = "  " * lvl
            md_lines.append(f"{indent}- {t_clean}")
            knowledge_points.append(t_clean)

        md_lines.append("")

    # 末尾附加知识点汇总
    md_lines.append("---")
    md_lines.append("")
    md_lines.append("## 知识点汇总")
    md_lines.append("")
    md_lines.append(f"> 共提取 {len(knowledge_points)} 条知识点")
    md_lines.append("")
    for i, kp in enumerate(knowledge_points, 1):
        md_lines.append(f"{i}. {kp}")

    return "\n".join(md_lines)


def main():
    if len(sys.argv) < 2:
        print("用法: python ppt2md.py <input.pptx> [output.md|--stdout]")
        sys.exit(1)

    pptx_path = sys.argv[1]
    if not os.path.exists(pptx_path):
        print(f"[错误] 文件不存在: {pptx_path}")
        sys.exit(1)

    md_content = ppt_to_markdown(pptx_path)

    # 确定输出目标
    if len(sys.argv) >= 3 and sys.argv[2] == "--stdout":
        print(md_content)
    else:
        if len(sys.argv) >= 3:
            out_path = sys.argv[2]
        else:
            base = os.path.splitext(pptx_path)[0]
            out_path = base + "_讲义.md"
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(md_content)
        print(f"[OK] 导出完成: {out_path}")
        # 统计
        kp_count = md_content.count("\n") 
        lines_count = len(md_content.splitlines())
        print(f"[INFO] 共 {lines_count} 行Markdown")


if __name__ == "__main__":
    main()